"""
NutriSL — Data Ingestion Script
================================
Converts the three raw source files into a single normalized SQLite database:

  1. FoodCompositionData_GI.xlsx  -> foods, food_nutrients, gi_gl_reference
  2. Sri_Lankan_DRI_Reference.xlsx -> dri_energy, dri_protein, dri_carbs_fibre,
                                       dri_fat_fatty_acids, dri_water,
                                       dri_vitamins, dri_minerals,
                                       dri_food_sources, dri_reference_body_weights
  3. Nutrition_goals.docx          -> disease_nutrition_goals, activity_factors,
                                       classification_cutoffs

Design choice: food nutrients are stored LONG (food_code, nutrient_code, value, unit)
rather than wide, because column counts vary per food group (64-116 cols) and this
lets the calculator/deficiency engine query any nutrient generically without
needing per-food-group schemas.

Run: python build_db.py
Produces: nutrisl.db in this same directory.
"""
import re
import sqlite3
from pathlib import Path

import openpyxl
import docx

HERE = Path(__file__).parent
UPLOADS = HERE / "source_files"
DB_PATH = HERE / "nutrisl.db"

FOOD_XLSX = UPLOADS / "FoodCompositionData_GI.xlsx"
DRI_XLSX = UPLOADS / "Sri_Lankan_DRI_Reference.xlsx"
GOALS_DOCX = UPLOADS / "Nutrition_goals__1_.docx"

FOOD_GROUP_SHEETS = [
    "SLA - Cereals & Grains", "SLB - Root Vegetables", "SLC - Legumes & Pulses",
    "SLD - Vegetables", "SLE - Fruits", "SLF - Fish & Aquatic",
    "SLG - Milk & Dairy Products", "SLH - Eggs, Poultry & Meat",
    "SLI - Nuts & Seeds", "SLJ - Oils & Fats", "SLK - Condiments & Spices",
    "SLL - Beverages",
]

# The 12 food-group sheets were digitized somewhat independently and use different
# nutrient codes for the same nutrient (e.g. Vitamin A is 'VIT-A' in some sheets,
# 'VITA' in others). Left unmerged, a diary combining foods from different groups
# would silently split one nutrient's total across two rows and undercount it in
# the deficiency check. Canonicalize at ingestion so downstream code only ever
# sees one code per real nutrient.
NUTRIENT_ALIASES = {
    "VIT-A": "VITA",
    "THIAMINE": "THIA",
    "RIBOFLAVIN": "RIBF",
    "NIACIN": "NIA",
    "VIT-D(ERGCAL)": "ERGCAL",
    "VIT-E": "VITE",
    "VIT-K": "VITK1",
    "FOLATE-SUM": "FOLSUM",
    "PANTOTHENIC": "PANTAC",
    "F18D2": "F18D2CN6",
    "F18D2C N6": "F18D2CN6",
    "F18D1": "F18D1CN9",
    "F18D1C": "F18D1CN9",
    "F18D1C N9": "F18D1CN9",
    "F18D3": "F18D3N3",
    "STIGSTR": "STGSTR",
    "a - Tocopherols": "TOCPHA",
    "a -Tocotrienols": "TOCTRA",
    "B- Tocopherols": "TOCPHB",
    "Y - Tocopherols": "TOCPHG",
    "6- Tocopherols": "TOCPHD",
}

NUTRIENT_CODE_RE = re.compile(r"\[([^\]]+)\]")
UNIT_RE = re.compile(r"\s([a-zA-ZµμΩ%°]+(?:/[a-zA-Z]+)?)$")


def parse_header_cell(cell_text):
    """'Energy\\n[ENERC] kcal' -> (name='Energy', code='ENERC', unit='kcal')"""
    if cell_text is None:
        return None, None, None
    text = str(cell_text)
    code_match = NUTRIENT_CODE_RE.search(text)
    code = code_match.group(1) if code_match else None
    name_part = text.split("\n")[0].strip()
    unit = None
    if code_match:
        after = text[code_match.end():].strip()
        unit = after if after else None
    return name_part, code, unit


# DATA QUALITY FIX: the 'Energy [ENERC] kcal' column across ALL 12 food-group sheets
# is actually in kilojoules, not kcal as labeled. Verified by cross-checking known real
# values: raw Barley=1311 -> /4.184=313.3kcal (correct); raw Canola oil=3700 -> /4.184=
# 884.3kcal (correct, matches the ~884kcal/100g standard for refined oils); raw Apple=214
# -> /4.184=51.1kcal (correct). Converting at ingestion so every downstream kcal figure
# (energy calculator, DRI comparison, disease goals) is actually in kcal.
KJ_TO_KCAL = 4.184


def clean_value(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v in ("-", "", "ND", "NE", "NA", "N/A"):
            return None
        try:
            return float(v)
        except ValueError:
            return None
    return float(v)


def build_schema(conn):
    conn.executescript(
        """
        DROP TABLE IF EXISTS foods;
        DROP TABLE IF EXISTS food_nutrients;
        DROP TABLE IF EXISTS nutrient_dictionary;
        DROP TABLE IF EXISTS gi_gl_reference;
        DROP TABLE IF EXISTS dri_energy;
        DROP TABLE IF EXISTS dri_protein;
        DROP TABLE IF EXISTS dri_carbs_fibre;
        DROP TABLE IF EXISTS dri_fat_fatty_acids;
        DROP TABLE IF EXISTS dri_water;
        DROP TABLE IF EXISTS dri_vitamins;
        DROP TABLE IF EXISTS dri_minerals;
        DROP TABLE IF EXISTS dri_food_sources;
        DROP TABLE IF EXISTS dri_reference_body_weights;
        DROP TABLE IF EXISTS disease_nutrition_goals;
        DROP TABLE IF EXISTS activity_factors;
        DROP TABLE IF EXISTS classification_cutoffs;

        CREATE TABLE foods (
            food_code TEXT PRIMARY KEY,
            food_name TEXT NOT NULL,
            food_group TEXT NOT NULL
        );

        CREATE TABLE nutrient_dictionary (
            nutrient_code TEXT PRIMARY KEY,
            nutrient_name TEXT,
            unit TEXT
        );

        CREATE TABLE food_nutrients (
            food_code TEXT NOT NULL,
            nutrient_code TEXT NOT NULL,
            value REAL,
            PRIMARY KEY (food_code, nutrient_code),
            FOREIGN KEY (food_code) REFERENCES foods(food_code)
        );

        CREATE TABLE gi_gl_reference (
            food_number INTEGER,
            item_description TEXT,
            gi_glucose100 REAL,
            gi_bread100 REAL,
            subjects TEXT,
            reference_food_time TEXT,
            serve_size_g REAL,
            avail_carb_g_per_serve REAL,
            gl_per_serve REAL,
            source_table TEXT
        );

        CREATE TABLE dri_energy (
            life_stage TEXT, age TEXT, sex TEXT, pal_activity TEXT,
            kcal_per_kg_day REAL, ar_energy_kcal_day REAL
        );

        CREATE TABLE dri_protein (
            life_stage TEXT, age TEXT, sex TEXT,
            ar_g_per_kg_bw_day REAL, rda_g_per_kg_bw_day REAL, ul_pe_ratio TEXT
        );

        CREATE TABLE dri_carbs_fibre (
            age_group TEXT, sex TEXT, total_carb_pct_e TEXT, dietary_fibre_g_day REAL
        );

        CREATE TABLE dri_fat_fatty_acids (
            age_group TEXT, total_fat_pct_e TEXT, sfa TEXT, la_pct_e TEXT,
            ala_pct_e TEXT, epa_dha_mg_day TEXT, tfa TEXT
        );

        CREATE TABLE dri_water (
            age_group TEXT, male_ai_l_day TEXT, female_ai_l_day TEXT
        );

        CREATE TABLE dri_vitamins (
            vitamin TEXT, unit TEXT, age_life_stage TEXT, sex TEXT,
            ar TEXT, rda TEXT, ai_ri TEXT, ul TEXT
        );

        CREATE TABLE dri_minerals (
            mineral TEXT, unit TEXT, age_life_stage TEXT, sex TEXT,
            ar TEXT, rda TEXT, ai_ri TEXT, ul TEXT
        );

        CREATE TABLE dri_food_sources (
            nutrient TEXT, unit TEXT, food_item TEXT,
            content_per_100g REAL, serving_size TEXT, content_per_serving REAL
        );

        CREATE TABLE dri_reference_body_weights (
            population_group TEXT, age_range TEXT, age_taken_as_reference TEXT,
            male_kg REAL, female_kg REAL, both_kg REAL
        );

        CREATE TABLE disease_nutrition_goals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            disease_condition TEXT NOT NULL,
            nutrition_goal_text TEXT NOT NULL
        );

        CREATE TABLE activity_factors (
            occupation TEXT, activity_level TEXT, factor_range TEXT
        );

        CREATE TABLE classification_cutoffs (
            category TEXT, cutoff_text TEXT
        );
        """
    )


def ingest_food_composition(conn):
    wb = openpyxl.load_workbook(FOOD_XLSX, read_only=True, data_only=True)
    nutrient_dict = {}
    food_rows = []
    fn_rows = []

    for sheet_name in FOOD_GROUP_SHEETS:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        header_row = rows[1]  # row with codes e.g. 'Energy\n[ENERC] kcal'
        data_rows = rows[2:]

        col_meta = []  # (col_idx, nutrient_code) for nutrient columns
        for idx, cell in enumerate(header_row):
            if idx < 2:
                continue  # food_code, food_name
            name, code, unit = parse_header_cell(cell)
            if code is None:
                continue
            code = NUTRIENT_ALIASES.get(code, code)
            if code not in nutrient_dict:
                nutrient_dict[code] = (name, unit)
            col_meta.append((idx, code))

        group_label = sheet_name.split(" - ", 1)[1] if " - " in sheet_name else sheet_name

        for row in data_rows:
            if not row or row[0] is None:
                continue
            food_code = str(row[0]).strip()
            food_name = str(row[1]).strip() if row[1] else ""
            food_rows.append((food_code, food_name, group_label))
            for idx, code in col_meta:
                val = clean_value(row[idx]) if idx < len(row) else None
                if val is not None:
                    if code == "ENERC":
                        val = round(val / KJ_TO_KCAL, 1)
                    fn_rows.append((food_code, code, val))

    conn.executemany(
        "INSERT OR REPLACE INTO nutrient_dictionary VALUES (?,?,?)",
        [(code, name, unit) for code, (name, unit) in nutrient_dict.items()],
    )
    conn.executemany("INSERT OR REPLACE INTO foods VALUES (?,?,?)", food_rows)
    conn.executemany("INSERT OR REPLACE INTO food_nutrients VALUES (?,?,?)", fn_rows)
    return len(food_rows), len(fn_rows), len(nutrient_dict)


def ingest_gi_gl(conn):
    wb = openpyxl.load_workbook(FOOD_XLSX, read_only=True, data_only=True)
    ws = wb["GI-GL Reference"]
    rows = list(ws.iter_rows(values_only=True))[2:]
    out = []
    for r in rows:
        if not r or r[0] is None:
            continue
        out.append((
            r[0], r[1], clean_value(r[2]), clean_value(r[3]), r[4], r[5],
            clean_value(r[7]), clean_value(r[8]), clean_value(r[9]), r[10],
        ))
    conn.executemany(
        "INSERT INTO gi_gl_reference VALUES (?,?,?,?,?,?,?,?,?,?)", out
    )
    return len(out)


def ingest_dri(conn):
    wb = openpyxl.load_workbook(DRI_XLSX, read_only=True, data_only=True)
    counts = {}

    def rows_after_header(sheet, n_header=2):
        rows = list(wb[sheet].iter_rows(values_only=True))
        out = []
        for r in rows[n_header:]:
            if not r or not any(c is not None for c in r):
                continue
            first = str(r[0]).strip() if r[0] is not None else ""
            if first.lower().startswith("source"):
                continue  # trailing footnote row, not data
            out.append(r)
        return out

    energy = rows_after_header("Energy")
    conn.executemany("INSERT INTO dri_energy VALUES (?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], r[3], clean_value(r[4]), clean_value(r[5])) for r in energy])
    counts["dri_energy"] = len(energy)

    protein = rows_after_header("Protein")
    conn.executemany("INSERT INTO dri_protein VALUES (?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], clean_value(r[3]), clean_value(r[4]), r[5]) for r in protein])
    counts["dri_protein"] = len(protein)

    carbs = rows_after_header("Carbs_Fibre")
    conn.executemany("INSERT INTO dri_carbs_fibre VALUES (?,?,?,?)",
                      [(r[0], r[1], r[2], clean_value(r[3])) for r in carbs])
    counts["dri_carbs_fibre"] = len(carbs)

    fat = rows_after_header("Fat_FattyAcids")
    conn.executemany("INSERT INTO dri_fat_fatty_acids VALUES (?,?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], r[3], r[4], r[5], r[6]) for r in fat])
    counts["dri_fat_fatty_acids"] = len(fat)

    water = rows_after_header("Water")
    conn.executemany("INSERT INTO dri_water VALUES (?,?,?)",
                      [(r[0], r[1], r[2]) for r in water])
    counts["dri_water"] = len(water)

    vitamins = rows_after_header("Vitamins")
    conn.executemany("INSERT INTO dri_vitamins VALUES (?,?,?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]) for r in vitamins])
    counts["dri_vitamins"] = len(vitamins)

    minerals = rows_after_header("Minerals")
    conn.executemany("INSERT INTO dri_minerals VALUES (?,?,?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]) for r in minerals])
    counts["dri_minerals"] = len(minerals)

    food_sources = rows_after_header("Food_Sources")
    conn.executemany("INSERT INTO dri_food_sources VALUES (?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], clean_value(r[3]), r[4], clean_value(r[5])) for r in food_sources])
    counts["dri_food_sources"] = len(food_sources)

    rbw = rows_after_header("Reference_BodyWeights")
    conn.executemany("INSERT INTO dri_reference_body_weights VALUES (?,?,?,?,?,?)",
                      [(r[0], r[1], r[2], clean_value(r[3]), clean_value(r[4]), clean_value(r[5])) for r in rbw])
    counts["dri_reference_body_weights"] = len(rbw)

    return counts


def ingest_goals_doc(conn):
    d = docx.Document(GOALS_DOCX)
    counts = {}

    # Table 0: disease -> nutrition goal
    t0 = d.tables[0]
    goal_rows = []
    for row in t0.rows[1:]:
        disease = row.cells[0].text.strip()
        goal = row.cells[1].text.strip()
        if disease and goal:
            goal_rows.append((disease, goal))
    conn.executemany(
        "INSERT INTO disease_nutrition_goals (disease_condition, nutrition_goal_text) VALUES (?,?)",
        goal_rows,
    )
    counts["disease_nutrition_goals"] = len(goal_rows)

    # Table 1: occupation -> activity level -> factor
    t1 = d.tables[1]
    act_rows = []
    for row in t1.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            act_rows.append(tuple(cells))
    conn.executemany("INSERT INTO activity_factors VALUES (?,?,?)", act_rows)
    counts["activity_factors"] = len(act_rows)

    # Table 2: BMI / waist / waist-hip classification cutoffs
    t2 = d.tables[2]
    cutoff_rows = []
    for row in t2.rows:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) >= 2 and cells[0]:
            cutoff_rows.append((cells[0], cells[1]))
    conn.executemany("INSERT INTO classification_cutoffs VALUES (?,?)", cutoff_rows)
    counts["classification_cutoffs"] = len(cutoff_rows)

    return counts


def main():
    if DB_PATH.exists():
        DB_PATH.unlink()
    conn = sqlite3.connect(DB_PATH)
    build_schema(conn)

    n_foods, n_fn, n_dict = ingest_food_composition(conn)
    n_gi = ingest_gi_gl(conn)
    dri_counts = ingest_dri(conn)
    goal_counts = ingest_goals_doc(conn)

    conn.commit()

    print(f"Foods ingested:            {n_foods}")
    print(f"Nutrient values ingested:  {n_fn}")
    print(f"Distinct nutrients:        {n_dict}")
    print(f"GI/GL reference rows:      {n_gi}")
    for k, v in dri_counts.items():
        print(f"{k:<30} {v}")
    for k, v in goal_counts.items():
        print(f"{k:<30} {v}")

    conn.close()
    print(f"\nDatabase written to: {DB_PATH}")


if __name__ == "__main__":
    main()
